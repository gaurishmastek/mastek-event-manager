"""The officer's verification step: check the employee's ID, tick who is present, approve or reject.

A shared pass admits the employee first, then any guests who arrive later, each person once.
"""

import pytest
from sqlalchemy import select

from app.modules.gate.models import CheckIn, EntryRejection, GuestEntry, ScanAttempt
from app.modules.gate.service import GateService
from app.modules.guests.models import Registration
from tests.gate.helpers import admit, decide, guest_ids, scan

PARTY = ["Ravi Patil", "Meera Patil", "Tia Patil"]  # adult, then two kids


@pytest.fixture()
def party_token(issue_pass, live_event, officer):
    return issue_pass(live_event.id, guests=PARTY, employee_id="MT-104")["qr_token"]


def entry_rows(db_session):
    return db_session.scalars(select(GuestEntry).order_by(GuestEntry.id)).all()


# --- approving the first visit -------------------------------------------------------


def test_approval_needs_the_employee_id_to_be_checked(client, live_event, party_token, db_session):
    ids = guest_ids(client, live_event.id, party_token)

    response = decide(client, live_event.id, party_token, guest_ids_entered=ids)

    assert response.status_code == 422
    assert "ID card" in response.json()["detail"]
    assert db_session.scalars(select(CheckIn)).all() == [] and entry_rows(db_session) == []
    assert db_session.scalars(select(Registration)).one().status == "VERIFIED"


def test_employee_alone_can_be_admitted_first(client, live_event, party_token, db_session):
    body = admit(client, live_event.id, party_token, guests=0).json()

    assert body["result"] == "admitted"
    assert body["people_entered"] == 1
    assert [m["entered"] for m in body["guest"]["members"]] == [False, False, False]
    assert body["guest"]["employee_entered"] is True
    assert entry_rows(db_session) == []


def test_some_guests_can_enter_with_the_employee(client, live_event, party_token, db_session):
    body = admit(client, live_event.id, party_token, guests=2, gate="North gate").json()

    assert body["people_entered"] == 3
    assert [m["entered"] for m in body["guest"]["members"]] == [True, True, False]
    rows = entry_rows(db_session)
    assert [(r.officer_id, r.gate) for r in rows] == [(42, "North gate")] * 2
    assert {r.check_in_id for r in rows} == {db_session.scalars(select(CheckIn)).one().id}


def test_guest_ids_must_belong_to_this_pass(client, issue_pass, live_event, party_token, db_session):
    other_pass = issue_pass(live_event.id, email="other@example.com", guests=["Sam"])
    other_ids = guest_ids(client, live_event.id, other_pass["qr_token"])

    for bad_ids in (other_ids, [987654]):
        response = decide(client, live_event.id, party_token, employee_id_checked=True, guest_ids_entered=bad_ids)
        assert response.status_code == 422
        assert "not on this pass" in response.json()["detail"]
    assert db_session.scalars(select(CheckIn)).all() == [] and entry_rows(db_session) == []


def test_guest_listed_twice_is_refused(client, live_event, party_token):
    first = guest_ids(client, live_event.id, party_token)[0]

    response = decide(client, live_event.id, party_token, employee_id_checked=True, guest_ids_entered=[first, first])

    assert response.status_code == 422


# --- late guests on the shared pass ------------------------------------------------------


def test_scan_after_the_employee_entered_lists_who_is_still_to_arrive(client, live_event, party_token):
    admit(client, live_event.id, party_token, guests=1, gate="North gate")

    body = scan(client, live_event.id, party_token, gate="South gate").json()

    assert body["result"] == "pending_guests"
    assert body["people_entered"] == 2
    assert body["gate"] == "North gate" and body["checked_in_at"]
    guest = body["guest"]
    assert guest["employee_entered"] is True and guest["employee_id"] == "MT-104"
    assert [(m["name"], m["entered"]) for m in guest["members"]] == [
        ("Ravi Patil", True),
        ("Meera Patil", False),
        ("Tia Patil", False),
    ]
    assert guest["members"][0]["entered_at"] and guest["members"][1]["entered_at"] is None


def test_late_guests_enter_on_the_same_pass_without_another_id_check(client, live_event, party_token, db_session):
    admit(client, live_event.id, party_token, guests=1, gate="North gate")
    ids = guest_ids(client, live_event.id, party_token)

    late = decide(client, live_event.id, party_token, gate="South gate", guest_ids_entered=[ids[1]]).json()

    assert late["result"] == "admitted"
    assert late["people_entered"] == 3
    assert [m["entered"] for m in late["guest"]["members"]] == [True, True, False]
    rows = entry_rows(db_session)
    assert [(r.registration_guest_id, r.gate) for r in rows] == [(ids[0], "North gate"), (ids[1], "South gate")]
    assert len(db_session.scalars(select(CheckIn)).all()) == 1  # the employee is not admitted twice

    last = decide(client, live_event.id, party_token, guest_ids_entered=[ids[2]]).json()
    assert last["people_entered"] == 4
    assert scan(client, live_event.id, party_token).json()["result"] == "already_checked_in"
    assert len(entry_rows(db_session)) == 3


def test_a_late_approval_must_tick_somebody(client, live_event, party_token, db_session):
    admit(client, live_event.id, party_token)

    response = decide(client, live_event.id, party_token, guest_ids_entered=[])

    assert response.status_code == 422
    assert entry_rows(db_session) == []


def test_a_guest_who_is_already_in_cannot_enter_twice(client, live_event, party_token, db_session):
    admit(client, live_event.id, party_token, guests=1)
    ids = guest_ids(client, live_event.id, party_token)

    body = decide(client, live_event.id, party_token, guest_ids_entered=[ids[0], ids[1]]).json()

    assert body["result"] == "guests_already_entered"
    assert [m["entered"] for m in body["guest"]["members"]] == [True, False, False]
    assert len(entry_rows(db_session)) == 1  # nothing recorded, not even the new guest


def test_two_gates_letting_in_the_same_guest_at_once_admit_them_once(
    client, live_event, party_token, db_session, monkeypatch
):
    """The other officer's entry commits after ours looked at the pass; the unique index refuses the second row."""
    admit(client, live_event.id, party_token)
    ids = guest_ids(client, live_event.id, party_token)
    check_in = db_session.scalars(select(CheckIn)).one()
    db_session.add(
        GuestEntry(
            check_in_id=check_in.id,
            registration_guest_id=ids[1],
            officer_id=99,
            gate="East gate",
            entered_at=check_in.checked_in_at,
        )
    )
    db_session.commit()
    real = GateService._guest_entries
    calls = []

    def stale_first_read(self, registration):
        calls.append(1)
        return {} if len(calls) == 1 else real(self, registration)

    monkeypatch.setattr(GateService, "_guest_entries", stale_first_read)

    body = decide(client, live_event.id, party_token, guest_ids_entered=[ids[1]]).json()

    assert body["result"] == "guests_already_entered"
    assert len(entry_rows(db_session)) == 1
    assert entry_rows(db_session)[0].officer_id == 99


# --- rejecting ------------------------------------------------------------------------


@pytest.mark.parametrize("reason", ["ID_MISMATCH", "ID_NOT_PRESENTED"])
def test_rejection_is_recorded_and_leaves_the_pass_usable(client, live_event, party_token, db_session, reason):
    response = decide(client, live_event.id, party_token, decision="reject", gate="Gate 2", reason=reason)

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["result"], body["guest"]) == ("rejected", None)
    rejection = db_session.scalars(select(EntryRejection)).one()
    assert (rejection.reason, rejection.note, rejection.officer_id, rejection.gate) == (reason, None, 42, "Gate 2")
    assert rejection.event_id == live_event.id and rejection.created_by == 42 and rejection.created_at
    assert db_session.scalars(select(CheckIn)).all() == []
    assert db_session.scalars(select(Registration)).one().status == "VERIFIED"
    assert db_session.scalars(select(ScanAttempt).order_by(ScanAttempt.id.desc())).first().result == "rejected"
    # The guest comes back with the right ID and the same pass is approved.
    assert admit(client, live_event.id, party_token).json()["result"] == "admitted"


def test_other_reason_needs_a_note(client, live_event, party_token, db_session):
    without = decide(client, live_event.id, party_token, decision="reject", reason="OTHER")
    blank = decide(client, live_event.id, party_token, decision="reject", reason="OTHER", note="   ")
    with_note = decide(
        client, live_event.id, party_token, decision="reject", reason="OTHER", note="  Badge belongs to someone else "
    )

    assert (without.status_code, blank.status_code, with_note.status_code) == (422, 422, 200)
    assert db_session.scalars(select(EntryRejection)).one().note == "Badge belongs to someone else"


def test_a_rejection_needs_a_valid_reason_and_cannot_admit_anyone(client, live_event, party_token, db_session):
    ids = guest_ids(client, live_event.id, party_token)
    bad_bodies = (
        {"decision": "reject"},
        {"decision": "reject", "reason": "BECAUSE"},
        {"decision": "reject", "reason": "ID_MISMATCH", "employee_id_checked": True},
        {"decision": "reject", "reason": "ID_MISMATCH", "guest_ids_entered": ids[:1]},
        {"decision": "approve", "employee_id_checked": True, "reason": "ID_MISMATCH"},
        {"decision": "approve", "employee_id_checked": True, "note": "hello there"},
        {"decision": "maybe"},
        {"decision": "reject", "reason": "OTHER", "note": "x" * 256},
    )

    for body in bad_bodies:
        response = client.post(f"/api/v1/gate/events/{live_event.id}/decision", json={"token": party_token, **body})
        assert response.status_code == 422, body
    assert db_session.scalars(select(EntryRejection)).all() == []
    assert db_session.scalars(select(CheckIn)).all() == []


def test_a_late_guest_can_be_rejected_without_changing_the_pass(client, live_event, party_token, db_session):
    admit(client, live_event.id, party_token)

    body = decide(
        client, live_event.id, party_token, decision="reject", reason="OTHER", note="Not on the family list"
    ).json()

    assert body["result"] == "rejected"
    assert db_session.scalars(select(EntryRejection)).one().note == "Not on the family list"
    assert scan(client, live_event.id, party_token).json()["result"] == "pending_guests"
    assert entry_rows(db_session) == []


def test_nothing_is_rejected_on_a_fully_used_pass(client, issue_pass, live_event, officer, db_session):
    token = issue_pass(live_event.id)["qr_token"]
    admit(client, live_event.id, token)

    body = decide(client, live_event.id, token, decision="reject", reason="ID_MISMATCH").json()

    assert body["result"] == "already_checked_in"
    assert db_session.scalars(select(EntryRejection)).all() == []


# --- what others see -------------------------------------------------------------------


def test_entries_show_how_many_of_the_party_are_in(client, live_event, party_token):
    admit(client, live_event.id, party_token, guests=1)
    ids = guest_ids(client, live_event.id, party_token)

    def listed():
        (item,) = client.get(f"/api/v1/gate/events/{live_event.id}/entries").json()["items"]
        return item

    assert (listed()["people_entered"], listed()["guest"]["party_size"]) == (2, 4)
    decide(client, live_event.id, party_token, guest_ids_entered=ids[1:])
    assert listed()["people_entered"] == 4


def test_admin_registration_list_and_export_count_people_entered(client, live_event, party_token, login_as, db_session):
    from openpyxl import load_workbook

    admit(client, live_event.id, party_token, guests=2)
    login_as("admin", user_id=1)

    listing = client.get(f"/api/v1/events/{live_event.id}/registrations").json()["items"]
    export = client.get(f"/api/v1/events/{live_event.id}/registrations/export.xlsx")

    assert [(r["people_entered"], r["party_size"]) for r in listing] == [(3, 4)]
    assert export.status_code == 200
    import io

    sheet = load_workbook(io.BytesIO(export.content)).active
    header, row = next(sheet.iter_rows(values_only=True)), list(sheet.iter_rows(values_only=True))[1]
    assert (header[-1], row[-1]) == ("People Entered", 3)


def test_guest_details_stay_locked_after_entry(client, live_event, party_token, login_as, db_session):
    admit(client, live_event.id, party_token, guests=1)
    registration = db_session.scalars(select(Registration)).one()
    login_as("admin", user_id=1)

    response = client.patch(
        f"/api/v1/events/{live_event.id}/registrations/{registration.public_id}",
        json={"employee_id": "MT-104", "employee_name": "Asha Patil", "adult_name": "Someone Else"},
    )

    assert response.status_code == 409, response.text
