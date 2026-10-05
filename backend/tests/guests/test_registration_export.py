"""GET /events/{id}/registrations/export.xlsx: complete, admin-only, PII-safe export."""

import uuid
from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.db.mixins import utcnow
from app.modules.guests.models import Registration, RegistrationStatus
from tests.conftest import auth_headers

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def url(event_id: int) -> str:
    return f"/api/v1/events/{event_id}/registrations/export.xlsx"


def sheet_from(response):
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    return workbook["Registrations"]


def test_admin_exports_every_registration_with_masked_contacts_and_ist_times(
    client, issue_pass, register, make_event, login_as, db_session
):
    event = make_event()
    issue_pass(
        event.id,
        email="asha.patil@example.com",
        name="Asha Patil",
        guests=["Ravi Patil", "Mira Patil"],
        employee_id="MT-104",
    )
    pending = register(event.id, email="formula@example.com", name="=2+2", employee_id="MT-200")
    assert pending.status_code == 202
    login_as("admin", user_id=7)

    response = client.get(url(event.id))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-disposition"] == f'attachment; filename="event-{event.id}-registrations.xlsx"'

    sheet = sheet_from(response)
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    records = [dict(zip(headers, row, strict=True)) for row in rows[1:]]
    assert len(records) == 2

    verified = next(row for row in records if row["Employee ID"] == "MT-104")
    assert verified["Employee Name"] == "Asha Patil"
    assert verified["Email (Masked)"] == "as•••@example.com"
    assert verified["Mobile (Masked)"] == "98•••••210"
    assert verified["Guest Names"] == "Ravi Patil, Mira Patil"
    assert verified["Adult Name"] == "Ravi Patil"
    assert verified["Kid Names"] == "Mira Patil"
    assert verified["Kid Ages"] == "8"
    assert verified["Party Size"] == 3
    assert verified["Status"] == "VERIFIED"
    assert verified["QR Issued"] == "Yes"

    database_row = next(row for row in db_session.query(Registration).all() if row.employee_id == "MT-104")
    expected_registered_at = database_row.created_at + timedelta(hours=5, minutes=30)
    assert abs(verified["Registered At (IST)"] - expected_registered_at) < timedelta(milliseconds=1)

    # Excel must treat formula-looking names as text, never executable formula cells.
    formula_record = next(row for row in records if row["Employee ID"] == "MT-200")
    assert formula_record["Employee Name"] == "'=2+2"
    formula_cell = next(cell for row in sheet.iter_rows() for cell in row if cell.value == "'=2+2")
    assert formula_cell.data_type != "f"

    exported_values = {str(value) for row in rows for value in row if value is not None}
    for secret in (
        "asha.patil@example.com",
        "+919876543210",
        database_row.email_hash,
        database_row.email_encrypted,
        database_row.mobile_hash,
        database_row.mobile_encrypted,
        database_row.qr_token_hash,
    ):
        assert secret not in exported_values


def test_export_is_not_limited_to_the_registration_page_size(client, make_event, login_as, db_session):
    event = make_event(capacity=100)
    now = utcnow()
    registrations = [
        Registration(
            public_id=str(uuid.uuid4()),
            event_id=event.id,
            employee_id=f"MT-{number:03d}",
            employee_id_normalized=f"MT-{number:03d}",
            employee_name=f"Employee {number:03d}",
            email_hash=f"hash-{number:03d}",
            email_encrypted=f"encrypted-{number:03d}",
            email_masked=f"em•••{number:03d}@example.com",
            status=RegistrationStatus.PENDING_OTP.value,
            consent_at=now,
        )
        for number in range(526)
    ]
    registrations[-1].deleted_at = now
    db_session.add_all(registrations)
    db_session.commit()
    login_as("admin", user_id=7)

    response = client.get(url(event.id))

    assert response.status_code == 200
    # One header plus 525 active rows: past both the 20-row UI page and the 500-row DB export batch.
    assert sum(1 for _ in sheet_from(response).iter_rows()) == 526


def test_export_requires_admin_and_an_existing_event(client, make_event, login_as, db_session):
    event = make_event()
    assert client.get(url(event.id)).status_code == 401

    login_as("security_officer", user_id=42)
    assert client.get(url(event.id)).status_code == 403

    login_as("admin", user_id=7)
    assert client.get(url(999)).status_code == 404

    event.deleted_at = utcnow()
    db_session.commit()
    assert client.get(url(event.id)).status_code == 404


def test_real_tokens_enforce_export_role(client, officer, admin_headers, make_event):
    event = make_event()
    assert client.get(url(event.id), headers=auth_headers(client, officer)).status_code == 403
    assert client.get(url(event.id), headers=admin_headers).status_code == 200
